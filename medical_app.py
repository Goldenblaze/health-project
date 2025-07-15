import os
os.environ["PROTOCOL_BUFFERS_PYTHON_IMPLEMENTATION"] = "python"
import base64
import re
import streamlit as st
import google.generativeai as genai
from fpdf import FPDF
import PyPDF2
from docx import Document
import tempfile

# --- Streamlit Config ---
st.set_page_config(
    page_title="Medical Helper",
    page_icon="🩺",
    layout="centered"
)

# --- Secure Configuration ---
try:
    api_key = st.secrets['GOOGLE_API_KEY'].strip('"\'')
    if not api_key.startswith('AI'):
        st.error("⚠️ Key format looks wrong. Gemini keys start with 'AI'")
        st.stop()
    genai.configure(api_key=api_key)
except Exception as e:
    st.error(f"Key Error: {str(e)}")
    st.stop()

# --- Improved File Text Extraction ---
def extract_text_from_file(uploaded_file):
    """Extract and clean text from files"""
    try:
        with tempfile.NamedTemporaryFile(delete=False) as tmp:
            tmp.write(uploaded_file.getvalue())
            tmp_path = tmp.name
        
        raw_text = ""
        if uploaded_file.type == "application/pdf":
            with open(tmp_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                raw_text = "\n".join([page.extract_text() for page in reader.pages])
        elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            doc = Document(tmp_path)
            raw_text = "\n".join([para.text for para in doc.paragraphs])
        else:
            raw_text = uploaded_file.getvalue().decode("utf-8")
        
        cleaned_text = re.sub(r'(?<!\n)\n(?!\n)', ' ', raw_text)
        cleaned_text = re.sub(r'\s+', ' ', cleaned_text)
        return cleaned_text.strip()
    
    except Exception as e:
        st.error(f"Error reading file: {str(e)}")
        return ""
    finally:
        try:
            if 'tmp_path' in locals() and os.path.exists(tmp_path):
                os.unlink(tmp_path)
        except Exception as e:
            st.warning(f"Couldn't delete temp file: {str(e)}")

# --- Enhanced Red Flag Detection ---
RED_FLAGS = {
    r"\bchest\s*pain\b": "🚨 CALL 911: May indicate heart attack!",
    r"\bdifficulty\s*breathing\b": "🚨 Seek emergency care immediately!",
    r"\bsudden\s*numbness\b": "🚨 Stroke warning - call emergency services!",
    r"\bsevere\s*bleeding\b": "🚨 Apply pressure and call for emergency help!",
    r"\bsuicidal\s*thoughts\b": "🚨 Contact emergency services or a crisis hotline immediately"
}

def check_red_flags(text):
    """Check for red flags even with line breaks between words"""
    if not text:
        return False
    text = text.lower()
    for flag_pattern, message in RED_FLAGS.items():
        if re.search(flag_pattern, text):
            st.error(message)
            st.error("Do not use this app for emergencies. Seek immediate medical attention.")
            return True
    return False

# --- Fixed PDF Generation ---
def create_medical_pdf(symptoms, guide, filename="medical_summary.pdf"):
    try:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", 'B', 16)
        pdf.cell(0, 10, "Your Medical Visit Summary", 0, 1, 'C')
        pdf.ln(5)
        
        pdf.set_font("Helvetica", size=11)
        safe_text = guide.encode('ascii', 'replace').decode('ascii')
        pdf.multi_cell(0, 7, f"Symptoms:\n{symptoms}".encode('ascii', 'replace').decode('ascii'))
        pdf.ln(5)
        pdf.multi_cell(0, 7, f"Guidance:\n{safe_text}")
        
        pdf.ln(10)
        pdf.set_font("Helvetica", 'I', 8)
        pdf.multi_cell(0, 5, "Disclaimer: This document is not medical advice. Always consult a healthcare professional.")
        
        pdf.output(filename)
        return filename
    except Exception as e:
        st.error(f"PDF generation failed: {str(e)}")
        return None

def show_pdf_preview(pdf_path):
    """Display PDF preview in Streamlit"""
    if pdf_path and os.path.exists(pdf_path):
        with open(pdf_path, "rb") as f:
            base64_pdf = base64.b64encode(f.read()).decode('utf-8')
        pdf_display = f'<iframe src="data:application/pdf;base64,{base64_pdf}" width="100%" height="600" type="application/pdf"></iframe>'
        st.markdown(pdf_display, unsafe_allow_html=True)

# --- Streamlit UI ---
def main():
    st.title("🩺 Medical Helper")
    st.warning("⚠️ NOT for emergencies. Always consult a real doctor.")

    # Health Literacy Slider
    st.sidebar.header("Settings")
    reading_level = st.sidebar.slider(
        "Health Literacy Level",
        min_value=1,
        max_value=5,
        value=3,
        help="1 = Very Simple, 5 = More Detailed"
    )
    
    grade_levels = {1: "5th grade", 2: "7th grade", 3: "9th grade", 
                   4: "11th grade", 5: "College level"}
    selected_level = grade_levels[reading_level]

    # Input Methods
    doctor_specialty = st.selectbox(
        "Doctor's Specialty (if known)",
        ["General Practitioner", "Cardiologist", "Neurologist", "Other"],
        index=0
    )

    uploaded_file = st.file_uploader("Upload medical notes (PDF/DOCX/TXT)", 
                                   type=["pdf", "docx", "txt"])
    symptoms = st.text_area("Or describe symptoms manually:")

    # Process uploaded file
    if uploaded_file:
        extracted_text = extract_text_from_file(uploaded_file)
        if extracted_text:
            st.success("📄 Extracted Text Preview:")
            st.text_area("Extracted Content", 
                        value=extracted_text, 
                        height=150, 
                        label_visibility="collapsed")
            symptoms = extracted_text
            
            if check_red_flags(symptoms):
                st.stop()

    elif symptoms:
        if check_red_flags(symptoms):
            st.stop()

    # Main Generation Button
    if st.button("✨ Generate Medical Guide"):
        if not symptoms:
            st.warning("Please describe symptoms or upload a file")
        else:
            with st.spinner("Generating your personalized guide..."):
                try:
                    prompt = f"""
                    You are a health assistant for patients seeing a {doctor_specialty}.
                    Given these symptoms: {symptoms}
                    
                    Create a guide at {selected_level} reading level:
                    1. Possible causes (plain language, no diagnosis)
                    2. 5 priority questions for the doctor
                    3. Appointment preparation tips
                    
                    Rules:
                    - Use bullet points
                    - For level 1-2: Max 1-2 syllable words
                    - For level 3-4: May include medical terms with explanations
                    - Never suggest treatments
                    - Use only standard ASCII characters
                    """
                    
                    model = genai.GenerativeModel('gemini-1.5-flash')
                    response = model.generate_content(prompt, stream=True)
                    
                    # Initialize streaming
                    stream_placeholder = st.empty()
                    full_response = []
                    
                    # Stream response chunks
                    for chunk in response:
                        full_response.append(chunk.text)
                        stream_placeholder.markdown(
                            f"## 📋 Your Guide ({selected_level} level)\n\n" + 
                            "".join(full_response)
                        )
                    
                    # Get final response
                    guide = "".join(full_response)
                    
                    # Generate PDF
                    pdf_path = create_medical_pdf(symptoms, guide)
                    if pdf_path:
                        st.markdown("---")
                        st.subheader("📄 PDF Preview")
                        show_pdf_preview(pdf_path)
                        
                        with open(pdf_path, "rb") as f:
                            st.download_button(
                                "💾 Download Full Summary",
                                f,
                                file_name="medical_summary.pdf",
                                mime="application/pdf"
                            )
                except Exception as e:
                    st.error(f"Generation failed: {str(e)}")

    st.markdown("---")
    st.caption("Built with ❤️ for patient education")

if __name__ == "__main__":
    main()